"""
Analyse des CV : extraction locale, pseudonymisation avant envoi à l'IA, plan B sans IA.

Des CV fictifs sont générés pendant les tests (PDF texte, Word, PDF scanné) ; aucun appel
réseau réel n'est fait.
Lancer :  python manage.py test recrutements
"""
import io
import json
import shutil
import tempfile
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings

from ia.client import IAIndisponible

from .cv import (
    FORMAT_NON_SUPPORTE, ILLISIBLE, OK, SCANNE, correspondance_competences, extraire_texte,
    liste_competences, pseudonymiser, recommandation_sans_ia,
)

NOM = "Joseph MBARGA NDZANA"
EMAIL = "joseph.mbarga@exemple.cm"
LIGNES_CV = [
    NOM,
    "Adresse : Quartier Bastos, Yaoundé",
    "Téléphone : +237 6 99 12 34 56",
    f"Email : {EMAIL}",
    "Date de naissance : 12/03/1995",
    "Situation matrimoniale : marié, 2 enfants",
    "Profil LinkedIn : https://linkedin.com/in/jmbarga",
    "",
    "EXPÉRIENCE PROFESSIONNELLE",
    "2019 - 2023 : Développeur web chez ACME Douala. Applications Python et Django,",
    "bases de données MySQL, API REST. Contact du responsable : 699123456.",
    "2017 - 2019 : Technicien informatique, maintenance du parc et support utilisateurs.",
    "",
    "COMPÉTENCES",
    "Python, Django, MySQL, Git, Linux, conduite de réunions",
    "Langues : français courant, anglais technique",
    "",
    "FORMATION",
    "Licence professionnelle en génie logiciel, Université de Yaoundé I",
]
COMPETENCES_OFFRE = "Python, Django ; MySQL\n• Docker, Anglais technique, Gestion de projet"


def cv_pdf(lignes=LIGNES_CV):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    for ligne in lignes:
        pdf.cell(0, 6, ligne, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def cv_docx(lignes=LIGNES_CV):
    import docx
    document = docx.Document()
    for ligne in lignes:
        document.add_paragraph(ligne)
    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


def cv_scanne():
    """PDF ne contenant qu'une image (comme un CV photocopié puis scanné)."""
    from fpdf import FPDF
    from PIL import Image
    image = io.BytesIO()
    Image.new("RGB", (400, 560), "white").save(image, format="PNG")
    image.seek(0)
    pdf = FPDF()
    pdf.add_page()
    pdf.image(image, x=10, y=10, w=180)
    return bytes(pdf.output())


class ExtractionTests(SimpleTestCase):

    def test_pdf_texte(self):
        texte, statut = extraire_texte(io.BytesIO(cv_pdf()), "cv.pdf")
        self.assertEqual(statut, OK)
        self.assertIn("Django", texte)

    def test_word_docx(self):
        texte, statut = extraire_texte(io.BytesIO(cv_docx()), "cv.docx")
        self.assertEqual(statut, OK)
        self.assertIn("Licence professionnelle", texte)

    def test_pdf_scanne_signale(self):
        self.assertEqual(extraire_texte(io.BytesIO(cv_scanne()), "cv.pdf")[1], SCANNE)

    def test_image_et_ancien_word(self):
        self.assertEqual(extraire_texte(io.BytesIO(b"x"), "cv.jpg")[1], SCANNE)
        self.assertEqual(extraire_texte(io.BytesIO(b"x"), "cv.doc")[1], FORMAT_NON_SUPPORTE)

    def test_fichier_endommage(self):
        self.assertEqual(extraire_texte(io.BytesIO(b"pas un pdf"), "cv.pdf")[1], ILLISIBLE)


class PseudonymisationTests(SimpleTestCase):

    def setUp(self):
        texte, _ = extraire_texte(io.BytesIO(cv_pdf()), "cv.pdf")
        self.resultat = pseudonymiser(texte, NOM, EMAIL)

    def test_identite_et_coordonnees_retirees(self):
        for donnee in ("Joseph", "MBARGA", "NDZANA", EMAIL, "99 12 34 56", "699123456",
                       "Bastos", "12/03/1995", "linkedin.com/in/jmbarga", "marié"):
            with self.subTest(donnee=donnee):
                self.assertNotIn(donnee, self.resultat)

    def test_competences_et_experience_conservees(self):
        for donnee in ("Python", "Django", "MySQL", "2019 - 2023", "Développeur web",
                       "Licence professionnelle", "anglais technique"):
            with self.subTest(donnee=donnee):
                self.assertIn(donnee, self.resultat)


class CorrespondanceTests(SimpleTestCase):

    def test_decoupage_des_competences(self):
        self.assertEqual(liste_competences(COMPETENCES_OFFRE),
                         ["Python", "Django", "MySQL", "Docker", "Anglais technique", "Gestion de projet"])

    def test_score_et_competences_manquantes(self):
        texte, _ = extraire_texte(io.BytesIO(cv_pdf()), "cv.pdf")
        r = correspondance_competences(texte, COMPETENCES_OFFRE)
        self.assertEqual(r["trouvees"], ["Python", "Django", "MySQL", "Anglais technique"])
        self.assertEqual(r["manquantes"], ["Docker", "Gestion de projet"])
        self.assertEqual(r["score"], 67)                       # 4 sur 6

    def test_insensible_aux_accents_et_a_la_casse(self):
        r = correspondance_competences("Maîtrise de la COMPTABILITÉ et de Sage", "comptabilite, Sage")
        self.assertEqual(r["score"], 100)

    def test_pas_de_rejet_automatique_sans_ia(self):
        self.assertEqual(recommandation_sans_ia(10), "A_EXAMINER")
        self.assertEqual(recommandation_sans_ia(60), "TELEPHONE")
        self.assertEqual(recommandation_sans_ia(80), "ENTRETIEN_RH")


MEDIA_TEMP = tempfile.mkdtemp(prefix="sirh-cv-")


@override_settings(MEDIA_ROOT=MEDIA_TEMP)
class AnalyseCandidatureTests(TestCase):

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA_TEMP, ignore_errors=True)

    def _candidature(self, contenu, nom_fichier="cv.pdf"):
        from .models import Candidature, OffreEmploi
        offre = OffreEmploi.objects.create(titre="Développeur Python", description="Applications RH",
                                           competences_requises=COMPETENCES_OFFRE)
        return Candidature.objects.create(offre=offre, nom_complet=NOM, email=EMAIL,
                                          cv=SimpleUploadedFile(nom_fichier, contenu))

    def _analyser(self, candidature, **mock_ia):
        from . import services
        with mock.patch.object(services, "appeler_ia", **mock_ia) as appel:
            services._run_analyse(candidature.pk)
        candidature.refresh_from_db()
        return appel, json.loads(candidature.analyse_cv_ia), candidature

    def test_ia_recoit_le_cv_pseudonymise(self):
        reponse = json.dumps({"score_correspondance": 72, "points_forts_cv": ["Django"], "manques": ["Docker"],
                              "synthese": "Bon profil.", "recommandation": "ENTRETIEN_RH",
                              "niveau_estime": "CONFIRME"})
        appel, analyse, candidature = self._analyser(self._candidature(cv_pdf()), return_value=reponse)
        prompt = appel.call_args.args[0]
        self.assertIn("Django", prompt)                         # le contenu du CV est bien lu
        for donnee in ("MBARGA", "Joseph", EMAIL, "99 12 34 56", "Bastos"):
            self.assertNotIn(donnee, prompt)                    # rien d'identifiant ne part
        self.assertEqual(candidature.score_cv_ia, 72)
        self.assertEqual(analyse["source"], "IA (CV pseudonymisé)")
        self.assertEqual(analyse["correspondance_competences"]["score"], 67)

    def test_reponse_ia_normalisee(self):
        # Le modèle écrit parfois la clé à l'anglaise ou sort des valeurs attendues
        reponse = json.dumps({"score_correspondance": "104", "recommendation": "prioritaire",
                              "niveau_estime": "expert", "synthese": "x"})
        _, analyse, candidature = self._analyser(self._candidature(cv_pdf()), return_value=reponse)
        self.assertEqual(analyse["recommandation"], "PRIORITAIRE")
        self.assertEqual(candidature.score_cv_ia, 100)
        self.assertIsNone(analyse["niveau_estime"])

    def test_sans_internet_score_par_correspondance(self):
        _, analyse, candidature = self._analyser(
            self._candidature(cv_docx(), "cv.docx"), side_effect=IAIndisponible("Pas de connexion au service d'IA : x"))
        self.assertEqual(candidature.score_cv_ia, 67)
        self.assertEqual(analyse["source"], "calcul sans IA")
        self.assertEqual(analyse["manques"], ["Docker", "Gestion de projet"])
        self.assertEqual(analyse["recommandation"], "TELEPHONE")
        self.assertIn("4 compétence(s) exigée(s) sur 6", analyse["synthese"])

    def test_cv_scanne_rien_n_est_envoye(self):
        appel, analyse, candidature = self._analyser(self._candidature(cv_scanne()), return_value="{}")
        appel.assert_not_called()
        self.assertIsNone(candidature.score_cv_ia)
        self.assertEqual(analyse["statut_cv"], SCANNE)
        self.assertIn("version PDF texte ou Word", analyse["synthese"])
