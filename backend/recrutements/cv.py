"""
Traitement des CV avant analyse — tout se fait localement, sur le serveur de l'entreprise.

1. extraire_texte()        : texte du CV (PDF texte, Word .docx, .txt) ; un PDF sans texte
                              (scanné) ou un format non pris en charge est signalé.
2. pseudonymiser()         : retire l'identité et les coordonnées du candidat avant tout envoi
                              à un service d'IA hébergé à l'étranger. Loi n° 2024/017 : le
                              transfert de données personnelles hors du Cameroun requiert une
                              autorisation préalable de l'APDP — seules les compétences et
                              l'expérience quittent donc le serveur.
3. correspondance_competences() : plan B sans IA — part des compétences exigées par l'offre
                              retrouvées dans le CV (score explicable, hors ligne).
"""
import os
import re
import unicodedata

LONGUEUR_MIN_TEXTE = 200        # en dessous : PDF image (scanné) ou CV quasi vide
LONGUEUR_MAX_ENVOI = 12000      # caractères transmis à l'IA

OK = "OK"
SCANNE = "SCANNE"
FORMAT_NON_SUPPORTE = "FORMAT_NON_SUPPORTE"
ILLISIBLE = "ILLISIBLE"

MESSAGES = {
    SCANNE: "Le CV ne contient pas de texte exploitable (document scanné ou image). "
            "Demander au candidat une version PDF texte ou Word.",
    FORMAT_NON_SUPPORTE: "Format de CV non pris en charge (PDF, Word .docx ou texte attendus).",
    ILLISIBLE: "Le fichier du CV est introuvable ou endommagé.",
}


# ── 1. Extraction ────────────────────────────────────────────────────────────

def extraire_texte(fichier, nom=None):
    """(texte, statut) depuis un fichier ouvert en binaire ou un FieldFile Django."""
    nom = (nom or getattr(fichier, "name", "") or "").lower()
    extension = os.path.splitext(nom)[1]
    try:
        if extension == ".pdf":
            from pypdf import PdfReader
            lecteur = PdfReader(fichier)
            if lecteur.is_encrypted:
                lecteur.decrypt("")
            texte = "\n".join((page.extract_text() or "") for page in lecteur.pages)
        elif extension == ".docx":
            import docx
            document = docx.Document(fichier)
            morceaux = [p.text for p in document.paragraphs]
            for tableau in document.tables:
                for ligne in tableau.rows:
                    morceaux.append(" | ".join(c.text for c in ligne.cells))
            texte = "\n".join(morceaux)
        elif extension == ".txt":
            brut = fichier.read()
            texte = brut.decode("utf-8") if isinstance(brut, bytes) else brut
        elif extension in (".jpg", ".jpeg", ".png", ".gif", ".webp", ".tif", ".tiff"):
            return "", SCANNE
        else:
            return "", FORMAT_NON_SUPPORTE
    except (FileNotFoundError, OSError):
        return "", ILLISIBLE
    except Exception:
        return "", ILLISIBLE
    texte = re.sub(r"[ \t]+", " ", texte).strip()
    if len(texte) < LONGUEUR_MIN_TEXTE:
        return texte, SCANNE
    return texte, OK


# ── 2. Pseudonymisation ──────────────────────────────────────────────────────

MASQUE_IDENTITE = "[INFORMATION PERSONNELLE MASQUÉE]"
RE_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
RE_LIEN = re.compile(r"(?:https?://|www\.)\S+|\b(?:linkedin|facebook|github)\.com/\S*", re.I)
# Numéros internationaux (+237…, 00 237…) et camerounais à 9 chiffres commençant par 6 ou 2
RE_TEL_INTERNATIONAL = re.compile(r"(?:\+|\b00)\d{1,3}[\s.\-]?(?:\(?\d\)?[\s.\-]?){6,12}\d")
RE_TEL_CAMEROUN = re.compile(r"\b[26](?:[\s.\-]?\d){8}\b")
# Lignes d'état civil, d'adresse ou d'identification
RE_LIGNE_IDENTITE = re.compile(
    r"^\s*[-•*]?\s*(adresse|domicile|r[ée]sidence|quartier|b\.?\s?p\.?\b|bo[îi]te postale|"
    r"n[ée]e? le|n[ée]e? à|date de naissance|lieu de naissance|[âa]ge\b|sexe\b|"
    r"nationalit[ée]|situation (matrimoniale|familiale)|[ée]tat civil|mari[ée]|c[ée]libataire|"
    r"religion|cni\b|carte (nationale )?d'identit[ée]|passeport|n° ?cnps|niu\b|"
    r"t[ée]l[ée]phone|t[ée]l\b|e-?mail|courriel)",
    re.I,
)


def pseudonymiser(texte, nom_complet="", email=""):
    """Texte du CV sans identité ni coordonnées ; compétences et expérience conservées."""
    lignes = []
    for ligne in texte.splitlines():
        lignes.append(MASQUE_IDENTITE if RE_LIGNE_IDENTITE.match(ligne) else ligne)
    texte = "\n".join(lignes)
    if email:
        texte = texte.replace(email, "[EMAIL]")
    texte = RE_EMAIL.sub("[EMAIL]", texte)
    texte = RE_LIEN.sub("[LIEN]", texte)
    texte = RE_TEL_INTERNATIONAL.sub("[TÉLÉPHONE]", texte)
    texte = RE_TEL_CAMEROUN.sub("[TÉLÉPHONE]", texte)
    for partie in re.split(r"[\s\-']+", nom_complet or ""):
        if len(partie) >= 2:
            texte = re.sub(rf"\b{re.escape(partie)}\b", "[CANDIDAT]", texte, flags=re.I)
    return texte[:LONGUEUR_MAX_ENVOI]


# ── 3. Plan B : correspondance des compétences ───────────────────────────────

def _normaliser(texte):
    sans_accents = unicodedata.normalize("NFKD", texte or "").encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", sans_accents.lower()).strip()


def liste_competences(competences_requises):
    """« Python, Django ; MySQL\n• Anglais » → ['Python', 'Django', 'MySQL', 'Anglais']."""
    elements = re.split(r"[,;\n•·]+", competences_requises or "")
    resultat = []
    for e in elements:
        e = e.strip(" -–*\t.")
        if e and e.lower() not in {c.lower() for c in resultat}:
            resultat.append(e)
    return resultat


def _present(competence, texte_normalise):
    c = _normaliser(competence)
    if not c:
        return False
    if re.search(rf"(?<![a-z0-9]){re.escape(c)}(?![a-z0-9])", texte_normalise):
        return True
    mots = [m for m in re.split(r"[^a-z0-9+#]+", c) if len(m) >= 3]
    return len(mots) > 1 and all(re.search(rf"(?<![a-z0-9]){re.escape(m)}(?![a-z0-9])", texte_normalise)
                                 for m in mots)


def correspondance_competences(texte_cv, competences_requises):
    """{'trouvees', 'manquantes', 'total', 'score'} — score = part des compétences retrouvées."""
    competences = liste_competences(competences_requises)
    texte_normalise = _normaliser(texte_cv)
    trouvees = [c for c in competences if _present(c, texte_normalise)]
    manquantes = [c for c in competences if c not in trouvees]
    score = round(len(trouvees) / len(competences) * 100) if competences else None
    return {"trouvees": trouvees, "manquantes": manquantes, "total": len(competences), "score": score}


def recommandation_sans_ia(score):
    """Jamais de rejet automatique sans IA ni regard humain."""
    if score is None:
        return "A_EXAMINER"
    if score >= 75:
        return "ENTRETIEN_RH"
    if score >= 50:
        return "TELEPHONE"
    return "A_EXAMINER"
